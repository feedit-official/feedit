from __future__ import annotations

from pathlib import Path

import joblib
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.pipeline import Pipeline
from sklearn.svm import LinearSVC


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_TRAINING_PATH = BASE_DIR / "noise_training_v1.xlsx"
DEFAULT_MODEL_PATH = BASE_DIR / "models" / "marketing_classifier_v1.joblib"

PROTECTED_TERMS = {
    "made",
    "자체제작",
}


def load_training_rows(path: Path) -> tuple[list[str], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["training_v1"]

    header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows())]
    column = {name: index for index, name in enumerate(header)}

    texts: list[str] = []
    labels: list[str] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        text = str(row[column["text"]] or "").strip()
        label = str(row[column["label"]] or "").strip().upper()

        if not text or label not in {"MARKETING", "KEEP"}:
            continue

        texts.append(text)
        labels.append(label)

    return texts, labels


def train(
    training_path: Path = DEFAULT_TRAINING_PATH,
    model_path: Path = DEFAULT_MODEL_PATH,
) -> Path:
    texts, labels = load_training_rows(training_path)

    marketing_count = labels.count("MARKETING")
    keep_count = labels.count("KEEP")

    if marketing_count < 20 or keep_count < 20:
        raise ValueError(
            f"학습 데이터 부족: MARKETING={marketing_count}, KEEP={keep_count}"
        )

    print("TRAINING:", training_path)
    print("MARKETING:", marketing_count)
    print("KEEP:", keep_count)
    print("TOTAL:", len(texts))

    model = Pipeline([
        (
            "tfidf",
            TfidfVectorizer(
                analyzer="char",
                ngram_range=(2, 5),
                lowercase=True,
                sublinear_tf=True,
                min_df=1,
            ),
        ),
        (
            "classifier",
            LinearSVC(
                class_weight="balanced",
                C=1.0,
            ),
        ),
    ])

    model.fit(texts, labels)

    bundle = {
        "version": 1,
        "model": model,
        "labels": ["KEEP", "MARKETING"],
        "auto_remove_margin": 1.0,
        "review_margin": 0.25,
        "protected_terms": sorted(PROTECTED_TERMS),
        "training_counts": {
            "MARKETING": marketing_count,
            "KEEP": keep_count,
        },
    }

    model_path.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(bundle, model_path)

    print("MODEL SAVED:", model_path)
    return model_path


if __name__ == "__main__":
    train()
